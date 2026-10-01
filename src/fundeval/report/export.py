"""报告导出：Markdown 与 Excel（正文第九、十部分）。"""

from __future__ import annotations

import tempfile
from pathlib import Path
from urllib.parse import quote

import pandas as pd

from fundeval.report.summary import SECTIONS, STYLE_CAVEAT, EvaluationReport, conclusion, format_value


def _md_escape(value) -> str:
    return str(value).replace("|", "\\|").replace("\n", " ")


def _md_table(df: pd.DataFrame) -> str:
    if df.empty:
        return "（无）"
    header = "| " + " | ".join(_md_escape(c) for c in df.columns) + " |"
    sep = "| " + " | ".join("---" for _ in df.columns) + " |"
    rows = ["| " + " | ".join(_md_escape(v) for v in row) + " |" for row in df.itertuples(index=False)]
    return "\n".join([header, sep, *rows])


def _regression_display(report: EvaluationReport) -> pd.DataFrame:
    tbl = report.regression_table()
    if tbl.empty:
        return tbl
    out = tbl.copy()
    for col in ("估计值", "标准误", "95% 下限", "95% 上限"):
        out[col] = out[col].map(lambda v: f"{v:.6f}")
    for col in ("t", "R²"):
        out[col] = out[col].map(lambda v: f"{v:.4f}")
    out["p"] = out["p"].map(lambda v: f"{v:.4f}")
    return out


def _robustness_display(report: EvaluationReport) -> pd.DataFrame:
    tbl = report.robustness_table()
    if tbl.empty:
        return tbl
    out = tbl.copy()
    for col in ("全样本估计", "剔除后估计"):
        out[col] = out[col].map(lambda v: f"{v:.6f}")
    for col in ("全样本 t", "剔除后 t"):
        out[col] = out[col].map(lambda v: f"{v:.4f}")
    for col in ("全样本 n", "剔除后 n"):
        out[col] = out[col].map(lambda v: "—" if pd.isna(v) else f"{int(v)}")
    return out


def _factor_display(report: EvaluationReport) -> pd.DataFrame:
    tbl = report.factor_table()
    if tbl.empty:
        return tbl
    out = tbl.copy()
    out["系数"] = out["系数"].map(lambda v: f"{v:.4f}")
    for col in ("t", "p"):
        out[col] = out[col].map(lambda v: f"{v:.4f}")
    for col in ("因子均值（每期）", "贡献（每期）", "贡献（算术年化）"):
        out[col] = out[col].map(lambda v: f"{v * 100:.4f}%")
    return out


def _factor_reconciliation_display(report: EvaluationReport) -> pd.DataFrame:
    rec = report.factor_reconciliation()
    out = rec.apply(lambda col: col.map(lambda v: f"{v * 100:.4f}%"))
    return out.rename_axis("项目").reset_index()


def _capacity_display(report: EvaluationReport) -> pd.DataFrame:
    cap = report.capacity_table()
    if cap.empty:
        return cap
    out = cap.copy()
    for col in ("交易金额", "日均成交额"):
        out[col] = out[col].map(lambda v: f"{v:,.0f}")
    out["所需参与率"] = out["所需参与率"].map(lambda v: "不适用" if pd.isna(v) else f"{v * 100:.2f}%")
    out["上限"] = out["上限"].map(lambda v: f"{v * 100:.2f}%")
    return out.rename_axis("资产").reset_index()


def _components_display(report: EvaluationReport) -> pd.DataFrame:
    out = report.benchmark_components.copy()
    out["权重"] = out["权重"].map(lambda v: "—" if pd.isna(v) else f"{v * 100:.0f}%")
    return out


def _rolling_style_display(report: EvaluationReport) -> pd.DataFrame:
    roll = report.style_rolling
    out = roll.weights.apply(lambda col: col.map(lambda v: f"{v * 100:.1f}%"))  # DataFrame.map 需 pandas 2.1
    out["R²"] = roll.r_squared.map(lambda v: f"{v:.3f}")
    out.index = out.index.strftime("%Y-%m-%d") if isinstance(out.index, pd.DatetimeIndex) else out.index
    return out.rename_axis("窗口末期").reset_index()


def charts_dir(path: str | Path) -> Path:
    """Markdown 报告的图片目录：与报告同目录的“<报告名>_files/”。"""
    path = Path(path)
    return path.parent / f"{path.stem}_files"


def markdown_images(chart_set, md_path: str | Path) -> list[str]:
    """图表的 Markdown 嵌入行（相对报告文件的路径，按 URL 规则转义空格等字符），附中文图注。"""
    from fundeval.report.charts import CHART_TITLES

    base = Path(md_path).parent
    lines = []
    for key, png in chart_set.paths.items():
        rel = Path(png).resolve().relative_to(base.resolve()).as_posix()
        lines += [f"![{CHART_TITLES.get(key, key)}]({quote(rel)})", ""]
    lines += [f"- 未生成“{CHART_TITLES.get(k, k)}”：{reason}" for k, reason in chart_set.skipped.items()]
    if chart_set.lang == "en":
        lines.append("- 未找到中文字体，图表使用英文标签。")
    return lines


def insert_chart_sheet(book, chart_set, sheet_name: str = "图表", scale: float = 0.6) -> None:
    """在 openpyxl 工作簿中新增“图表”sheet，依次插入 PNG（按 scale 缩放），每张图上方写图注。

    图片在保存工作簿时才读取，调用方须保证 PNG 文件在保存前仍存在。
    """
    from openpyxl.drawing.image import Image

    from fundeval.report.charts import CHART_TITLES

    ws = book.create_sheet(sheet_name)
    row = 1
    for key, png in chart_set.paths.items():
        ws.cell(row=row, column=1, value=CHART_TITLES.get(key, key))
        img = Image(str(png))
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws.add_image(img, f"A{row + 1}")
        row += int(img.height / 20) + 4  # 默认行高 20 像素
    for key, reason in chart_set.skipped.items():
        ws.cell(row=row, column=1, value=f"未生成“{CHART_TITLES.get(key, key)}”：{reason}")
        row += 1
    if chart_set.lang == "en":
        ws.cell(row=row, column=1, value="未找到中文字体，图表使用英文标签。")


def to_markdown(report: EvaluationReport, path: str | Path | None = None, *, charts: bool = False) -> str:
    """生成 Markdown 报告：标题、口径、基准解析（用合同基准时）、各维度指标表（含收益来源、成本与容量）、回归系数表、
    多因子暴露与贡献、容量检查、滚动风格权重、稳健性检验、数据质量、结论与未完成检验。

    ``path`` 给出时同时写入文件（UTF-8）。返回 Markdown 文本。
    回归系数为每期值（与输入收益同频），未换算为百分数。
    ``charts=True`` 时（需 ``fundeval[plot]`` 与 ``path``）把 PNG 图表（report.charts.report_charts）保存到
    “<报告名>_files/”目录，在“图表”一节用相对路径嵌入。
    """
    if charts and path is None:
        raise ValueError("charts=True 须给出 path：图片保存到报告旁的“<报告名>_files/”目录")
    chart_lines: list[str] = []
    if charts:
        from fundeval.report.charts import report_charts

        chart_lines = markdown_images(report_charts(report, charts_dir(path)), path)
    parts = [f"# {report.title}", "", "## 口径", "", _md_table(pd.DataFrame(list(report.scope.items()), columns=["项目", "内容"]))]
    if report.benchmark_components is not None and len(report.benchmark_components):
        parts += ["", "## 基准解析", "", "合同业绩比较基准逐项解析；权重为合同权重，每期再平衡。", "",
                  _md_table(_components_display(report))]
    for name in SECTIONS:
        section = report.section(name)
        if section.empty:
            continue
        parts += ["", f"## {name}", "", _md_table(section)]
    reg = _regression_display(report)
    if not reg.empty:
        parts += ["", "## 回归系数", "", "系数为每期值，与输入收益同频。", "", _md_table(reg)]
    if report.factor is not None:
        intro = "贡献 = 因子系数 × 因子均值，算术年化乘以 K；各因子贡献、Alpha 与残差之和等于组合平均超额收益。"
        if report.factor.caveat:
            intro += f"{report.factor.caveat}。"
        parts += [
            "", "## 多因子暴露与贡献", "", intro, "", _md_table(_factor_display(report)),
            "", "收益对账：", "", _md_table(_factor_reconciliation_display(report)),
        ]
    if report.capacity is not None:
        parts += ["", "## 容量检查", "", "所需参与率 = 交易金额 / (日均成交额 × 可用天数)。", "",
                  _md_table(_capacity_display(report))]
    if report.style_rolling is not None:
        parts += [
            "", f"## 滚动风格权重（窗口 {report.style_rolling.window} 期）", "",
            f"按窗口末期排列。{STYLE_CAVEAT}；权重变化只提示进一步检查持仓与投资授权。", "",
            _md_table(_rolling_style_display(report)),
        ]
    if report.robustness is not None:
        parts += ["", "## 稳健性：剔除异常期", "", report.robustness.summary()]
        rob = _robustness_display(report)
        if not rob.empty:
            parts += ["", "异常期取自数据质量报告（组合、基准或市场任一被标记即剔除）；系数为每期值。", "", _md_table(rob)]
    if chart_lines:
        parts += ["", "## 图表", "", *chart_lines]
    parts += ["", "## 数据质量", "", _md_table(report.quality.summary())]
    issues = report.quality.issues()
    if issues:
        parts += ["", *[f"- {i}" for i in issues]]
    if report.notes:
        parts += ["", "## 附注", "", *[f"- {n}" for n in report.notes]]
    parts += ["", "## 结论", "", conclusion(report)]
    parts += ["", "## 未完成的检验", "", _md_table(report.not_done), ""]
    text = "\n".join(parts)
    if path is not None:
        Path(path).write_text(text, encoding="utf-8")
    return text


def _style_sheet(report: EvaluationReport, writer) -> None:
    st = report.style
    k = report.periods_per_year
    summary = pd.DataFrame(
        [
            ("R²", st.r_squared),
            ("残差均值（每期）", st.residual_mean),
            ("残差均值（算术年化）", st.annualized_residual_mean(k)),
            ("残差年化波动", st.annualized_residual_volatility(k)),
            ("样本期数 n", st.n),
            ("目标函数", "最小化残差方差" if st.objective == "variance" else "最小化残差平方和"),
            ("说明", STYLE_CAVEAT),
        ],
        columns=["项目", "数值"],
    )
    st.table().to_excel(writer, sheet_name="风格分析", index=False)
    row = len(st.weights) + 2
    summary.to_excel(writer, sheet_name="风格分析", index=False, startrow=row)
    if report.style_rolling is not None:
        roll = report.style_rolling.table()
        roll.index = roll.index.strftime("%Y-%m-%d") if isinstance(roll.index, pd.DatetimeIndex) else roll.index
        roll.index.name = f"窗口末期（窗口 {report.style_rolling.window} 期）"
        roll.to_excel(writer, sheet_name="风格分析", startrow=row + len(summary) + 3)


def to_excel(report: EvaluationReport, path: str | Path, *, charts: bool = False) -> Path:
    """写入 Excel（openpyxl），分 sheet：口径、基准解析（用合同基准时）、指标、回归、多因子、稳健性、风格分析、成本与容量、期间收益、数据质量。

    多因子 sheet 为因子暴露表，其下为收益对账（每期与算术年化）；成本与容量 sheet 为容量检查表
    （给出 costs['capacity'] 时）。

    风格分析 sheet 先列全样本权重、R² 与残差统计，其下为滚动权重（给出 style_window 时）；
    稳健性 sheet 为剔除异常期前后的关键系数与说明。

    指标 sheet 保存原始数值（小数）与单位，另附格式化后的显示值；期间收益 sheet 为对齐后的
    单期收益、主动收益与财富指数。需要安装 ``fundeval[excel]``。

    ``charts=True`` 时（需 ``fundeval[plot]``）另加“图表”sheet，用 openpyxl 插入 PNG 图表。
    """
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise ImportError('导出 Excel 需要 openpyxl：pip install "fundeval[excel]"') from exc
    path = Path(path)
    scope = pd.DataFrame(list(report.scope.items()), columns=["项目", "内容"])
    metrics = report.metrics.reset_index().rename(
        columns={"section": "维度", "key": "键", "label": "指标", "value": "数值", "unit": "单位", "note": "口径说明"}
    )
    metrics["显示值"] = [format_value(v, u, 4) if u in ("pct", "pp", "ratio") else format_value(v, u)
                       for v, u in zip(report.metrics["value"], report.metrics["unit"])]
    metrics["数值"] = [
        v if isinstance(v, (int, float)) or v is None else (f"{v:%Y-%m-%d}" if isinstance(v, pd.Timestamp) else str(v))
        for v in metrics["数值"]
    ]
    regression = report.regression_table()
    data = report.data.copy()
    data.index = data.index.strftime("%Y-%m-%d")
    data.index.name = "日期"
    quality = report.quality.summary()
    issues = pd.DataFrame({"需复核的问题": report.quality.issues()})
    conclusion_df = pd.DataFrame({"结论": conclusion(report).split("\n\n")})
    with tempfile.TemporaryDirectory() as tmp:
        chart_set = None
        if charts:
            from fundeval.report.charts import report_charts

            chart_set = report_charts(report, tmp)
        _write_excel(report, path, scope, metrics, regression, data, quality, issues, conclusion_df, chart_set)
    return path


def _write_excel(report, path, scope, metrics, regression, data, quality, issues, conclusion_df, chart_set) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        scope.to_excel(writer, sheet_name="口径", index=False)
        if report.benchmark_components is not None and len(report.benchmark_components):
            report.benchmark_components.to_excel(writer, sheet_name="基准解析", index=False)
        metrics.to_excel(writer, sheet_name="指标", index=False)
        regression.to_excel(writer, sheet_name="回归", index=False)
        if report.factor is not None:
            ftbl = report.factor_table()
            ftbl.to_excel(writer, sheet_name="多因子", index=False)
            report.factor_reconciliation().rename_axis("项目").to_excel(
                writer, sheet_name="多因子", startrow=len(ftbl) + 2
            )
            if report.factor.caveat:
                pd.DataFrame({"说明": [report.factor.caveat]}).to_excel(
                    writer, sheet_name="多因子", index=False, startrow=len(ftbl) + len(report.factor_reconciliation()) + 5
                )
        if report.robustness is not None:
            pd.DataFrame({"说明": [report.robustness.summary()]}).to_excel(writer, sheet_name="稳健性", index=False)
            rob = report.robustness_table()
            if not rob.empty:
                rob.to_excel(writer, sheet_name="稳健性", index=False, startrow=3)
        if report.style is not None:
            _style_sheet(report, writer)
        if report.capacity is not None:
            report.capacity.rename_axis("资产").to_excel(writer, sheet_name="成本与容量")
        data.to_excel(writer, sheet_name="期间收益")
        quality.to_excel(writer, sheet_name="数据质量", index=False)
        issues.to_excel(writer, sheet_name="数据质量", index=False, startrow=len(quality) + 2)
        conclusion_df.to_excel(writer, sheet_name="结论", index=False)
        report.not_done.to_excel(writer, sheet_name="未完成的检验", index=False)
        if chart_set is not None:
            insert_chart_sheet(writer.book, chart_set)
